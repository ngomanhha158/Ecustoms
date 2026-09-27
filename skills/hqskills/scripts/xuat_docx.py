# -*- coding: utf-8 -*-
"""Xuất văn bản hành chính ra .docx theo thể thức Nghị định 30/2020/NĐ-CP (Phụ lục I).

Đầu vào là một tệp JSON (Claude soạn nội dung, script này chỉ lo thể thức):

{
  "co_quan_chu_quan": "CHI CỤC HẢI QUAN X",     // có thể bỏ trống
  "co_quan_ban_hanh": "ĐỘI NGHIỆP VỤ",
  "so_ky_hieu": "…/TTr-ĐNV",                     // chưa có số thì để "…"
  "dia_danh": "Hải Phòng",
  "ngay": "",                                     // "05/09/2026"; trống = "ngày … tháng … năm …"
  "loai": "TỜ TRÌNH",                             // CÔNG VĂN thì trích yếu in "V/v …" dưới số ký hiệu
  "trich_yeu": "Về việc …",
  "kinh_gui": ["Chi cục Hải quan X"],
  "noi_dung": ["Đoạn mở đầu …", "I. THÔNG TIN LÔ HÀNG", "| Cột 1 | Cột 2 |", "| a | b |", "1. …"],
  "noi_nhan": ["Như trên", "Lưu: VT, Đội NV"],
  "quyen_han": "KT. ĐỘI TRƯỞNG",                  // có thể bỏ trống
  "chuc_vu": "PHÓ ĐỘI TRƯỞNG",
  "nguoi_ky": "Lê Văn …"
}

Trong noi_dung: dòng bắt đầu bằng "|" là hàng bảng (dòng "|---|" bị bỏ); dòng "I.", "II." … in đậm;
**chữ đậm** và *chữ nghiêng* được giữ. Chạy:

    python scripts/xuat_docx.py van_ban.json --ra to_trinh.docx
"""
import argparse
import json
import re
import sys

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
except ImportError:
    raise SystemExit("Thiếu thư viện python-docx — cài bằng: pip install python-docx")

PHONG = "Times New Roman"
CO_CHU = 14          # nội dung: cỡ 13–14
LA_MA = re.compile(r"^(?:[IVX]+)\.\s")
QUOC_HIEU = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"
_SAU_PBDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
             "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
             "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
             "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
             "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
TIEU_NGU = "Độc lập - Tự do - Hạnh phúc"


def _font(run, co, dam=False, nghieng=False):
    run.font.name = PHONG
    run._element.rPr.rFonts.set(qn("w:eastAsia"), PHONG)
    run.font.size = Pt(co)
    run.bold, run.italic = dam, nghieng
    return run


def _doan(noi, text="", co=CO_CHU, dam=False, nghieng=False, can=WD_ALIGN_PARAGRAPH.CENTER, sau=0, truoc=0):
    p = noi.add_paragraph()
    p.alignment = can
    p.paragraph_format.space_after = Pt(sau)
    p.paragraph_format.space_before = Pt(truoc)
    if text:
        _font(p.add_run(text), co, dam, nghieng)
    return p


def _chen_chu(p, text, co=CO_CHU, dam=False):
    """Giữ **đậm** và *nghiêng* trong một dòng."""
    for phan in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not phan:
            continue
        if phan.startswith("**"):
            _font(p.add_run(phan[2:-2]), co, True)
        elif phan.startswith("*"):
            _font(p.add_run(phan[1:-1]), co, dam, True)
        else:
            _font(p.add_run(phan), co, dam)


def _duong_ke(o, rong_cm, rong_o_cm):
    """Đường kẻ ngang nét liền, canh giữa trong ô, dài rong_cm."""
    p = o.add_paragraph()
    le = max((rong_o_cm - rong_cm) / 2, 0)
    p.paragraph_format.left_indent, p.paragraph_format.right_indent = Cm(le), Cm(le)
    p.paragraph_format.space_after = Pt(0)
    bd = OxmlElement("w:pBdr")
    duoi = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
        duoi.set(qn(k), v)
    bd.append(duoi)
    # OOXML bắt buộc thứ tự con của pPr: pBdr phải đứng trước spacing/ind/jc…, sai thứ tự Word báo lỗi tệp
    p._p.get_or_add_pPr().insert_element_before(bd, *_SAU_PBDR)
    _font(p.add_run(), 2)


def _bo_vien(bang):
    tblPr = bang._tbl.tblPr
    vien = OxmlElement("w:tblBorders")
    for canh in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{canh}")
        e.set(qn("w:val"), "nil")
        vien.append(e)
    tblPr.insert_element_before(vien, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
                                "w:tblCaption", "w:tblDescription", "w:tblPrChange")


def _ngay(vb):
    d = (vb.get("ngay") or "").strip()
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", d)
    if m:
        dd, mm, yy = int(m[1]), int(m[2]), m[3]
        # NĐ 30: ngày dưới 10 và tháng 1, 2 thêm số 0
        thang = f"{mm:02d}" if mm < 3 else str(mm)
        return f"ngày {dd:02d} tháng {thang} năm {yy}"
    return "ngày … tháng … năm …"


def _dau_van_ban(doc, vb):
    bang = doc.add_table(rows=1, cols=2)
    _bo_vien(bang)
    rong = (6.5, 9.5)
    trai, phai = bang.rows[0].cells
    for o, r in zip((trai, phai), rong):
        o.width = Cm(r)
    trai.paragraphs[0].text, phai.paragraphs[0].text = "", ""
    trai._tc.remove(trai.paragraphs[0]._p)
    phai._tc.remove(phai.paragraphs[0]._p)

    if vb.get("co_quan_chu_quan"):
        _doan(trai, vb["co_quan_chu_quan"].upper(), 13)
    _doan(trai, vb.get("co_quan_ban_hanh", "…").upper(), 13, dam=True)
    _duong_ke(trai, 2.2, rong[0])
    _doan(trai, f"Số: {vb.get('so_ky_hieu') or '…'}", 13, truoc=6)
    if vb.get("loai", "").upper() == "CÔNG VĂN" and vb.get("trich_yeu"):
        yeu = vb["trich_yeu"]
        _doan(trai, yeu if yeu.startswith("V/v") else f"V/v {yeu}", 12)

    _doan(phai, QUOC_HIEU, 13, dam=True)
    _doan(phai, TIEU_NGU, 14, dam=True)
    _duong_ke(phai, 5.8, rong[1])
    _doan(phai, f"{vb.get('dia_danh') or '…'}, {_ngay(vb)}", 14, nghieng=True, truoc=6)


def _bang(doc, dong):
    hang = [[c.strip() for c in d.strip().strip("|").split("|")] for d in dong
            if not re.fullmatch(r"\|?[\s:\-|]+\|?", d.strip())]
    so_cot = max(len(h) for h in hang)
    bang = doc.add_table(rows=len(hang), cols=so_cot)
    bang.style = "Table Grid"
    bang.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(hang):
        for j in range(so_cot):
            o = bang.cell(i, j)
            p = o.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            _chen_chu(p, h[j] if j < len(h) else "", 13, dam=(i == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def _noi_dung(doc, dong):
    i = 0
    while i < len(dong):
        d = dong[i].rstrip()
        if d.lstrip().startswith("|"):
            j = i
            while j < len(dong) and dong[j].lstrip().startswith("|"):
                j += 1
            _bang(doc, dong[i:j])
            i = j
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = p.paragraph_format
        pf.first_line_indent = Cm(1.27)
        pf.space_after = Pt(6)
        pf.line_spacing = 1.15
        _chen_chu(p, d, CO_CHU, dam=bool(LA_MA.match(d)))
        i += 1


def _cuoi_van_ban(doc, vb):
    bang = doc.add_table(rows=1, cols=2)
    _bo_vien(bang)
    trai, phai = bang.rows[0].cells
    trai.width, phai.width = Cm(8), Cm(8)
    trai._tc.remove(trai.paragraphs[0]._p)
    phai._tc.remove(phai.paragraphs[0]._p)

    p = _doan(trai, can=WD_ALIGN_PARAGRAPH.LEFT)
    _font(p.add_run("Nơi nhận:"), 12, dam=True, nghieng=True)
    ds = vb.get("noi_nhan") or ["Như trên", "Lưu: VT"]
    for k, x in enumerate(ds):
        _doan(trai, f"- {x}{'.' if k == len(ds) - 1 else ';'}", 11, can=WD_ALIGN_PARAGRAPH.LEFT)

    if vb.get("quyen_han"):
        _doan(phai, vb["quyen_han"].upper(), 14, dam=True)
    _doan(phai, (vb.get("chuc_vu") or "…").upper(), 14, dam=True)
    _doan(phai, "(Chữ ký, dấu)", 13, nghieng=True, sau=60)
    _doan(phai, vb.get("nguoi_ky") or "…", 14, dam=True)


def xuat(vb, ra):
    doc = Document()
    for s in doc.sections:   # A4, lề theo NĐ 30: trên/dưới 2–2,5 cm, trái 3–3,5 cm, phải 1,5–2 cm
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = Cm(2), Cm(2), Cm(3), Cm(1.5)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = PHONG, Pt(CO_CHU)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), PHONG)

    _dau_van_ban(doc, vb)
    if vb.get("loai", "").upper() != "CÔNG VĂN":
        _doan(doc, vb.get("loai", "").upper(), 14, dam=True, truoc=18)
        if vb.get("trich_yeu"):
            _doan(doc, vb["trich_yeu"], 14, dam=True)
        _duong_ke(doc, 3.5, 16.5)
    kg = vb.get("kinh_gui") or []
    if kg:
        p = _doan(doc, can=WD_ALIGN_PARAGRAPH.CENTER, truoc=12, sau=6)
        _font(p.add_run("Kính gửi: " + ("" if len(kg) > 1 else kg[0] + ".")), 14)
        if len(kg) > 1:
            for k, x in enumerate(kg):
                _doan(doc, f"- {x}{'.' if k == len(kg) - 1 else ';'}", 14, sau=0)
    _noi_dung(doc, vb.get("noi_dung") or [])
    _cuoi_van_ban(doc, vb)
    doc.save(ra)
    return ra


def main():
    ap = argparse.ArgumentParser(description="Xuất văn bản hành chính .docx theo NĐ 30/2020/NĐ-CP")
    ap.add_argument("tep_json", help="Tệp JSON nội dung văn bản ('-' = đọc từ stdin)")
    ap.add_argument("--ra", required=True, help="Đường dẫn tệp .docx ghi ra")
    a = ap.parse_args()
    vb = json.load(sys.stdin if a.tep_json == "-" else open(a.tep_json, encoding="utf-8"))
    print(f"Đã ghi {xuat(vb, a.ra)}")


if __name__ == "__main__":
    main()
