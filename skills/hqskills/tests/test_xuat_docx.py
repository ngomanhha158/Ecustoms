# -*- coding: utf-8 -*-
"""Kiểm xuat_docx.py: thể thức ngày tháng NĐ 30 và tệp .docx mở lại được, đúng thứ tự phần tử OOXML."""
import json
import os
import re
import sys
import tempfile
import unittest
import zipfile

GOC = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(GOC, "scripts"))
try:
    import xuat_docx
except SystemExit:          # máy chưa cài python-docx
    xuat_docx = None


@unittest.skipIf(xuat_docx is None, "chưa cài python-docx")
class XuatDocx(unittest.TestCase):
    def test_ngay_thang_nd30(self):
        # ngày dưới 10 và tháng 1, 2 thêm số 0; tháng 3–12 giữ nguyên
        self.assertEqual(xuat_docx._ngay({"ngay": "5/9/2026"}), "ngày 05 tháng 9 năm 2026")
        self.assertEqual(xuat_docx._ngay({"ngay": "15/2/2026"}), "ngày 15 tháng 02 năm 2026")
        self.assertEqual(xuat_docx._ngay({"ngay": "01/12/2026"}), "ngày 01 tháng 12 năm 2026")
        self.assertEqual(xuat_docx._ngay({}), "ngày … tháng … năm …")

    def test_xuat_mau_to_trinh(self):
        from docx import Document
        vb = json.load(open(os.path.join(GOC, "mau", "to_trinh_mau.json"), encoding="utf-8"))
        with tempfile.TemporaryDirectory() as d:
            ra = xuat_docx.xuat(vb, os.path.join(d, "t.docx"))
            doc = Document(ra)
            chu = "\n".join(p.text for p in doc.paragraphs)
            self.assertIn("TỜ TRÌNH", chu)
            self.assertIn("Kính gửi: Chi cục Hải quan X.", chu)
            dau = [p.text for c in doc.tables[0].rows[0].cells for p in c.paragraphs]
            self.assertIn("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", dau)
            self.assertIn("…, ngày 05 tháng 9 năm 2026", dau)
            xml = zipfile.ZipFile(ra).read("word/document.xml").decode("utf-8")
        # pBdr phải đứng trước spacing/ind/jc trong pPr
        for ppr in re.findall(r"<w:pPr>(.*?)</w:pPr>", xml):
            if "<w:pBdr>" in ppr:
                for sau in ("<w:spacing", "<w:ind", "<w:jc"):
                    if sau in ppr:
                        self.assertLess(ppr.index("<w:pBdr>"), ppr.index(sau))


if __name__ == "__main__":
    unittest.main()
