#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Nhập dữ liệu từ file Biểu thuế XNK gốc dạng "BT2026" (đúng layout thật
của Bộ Tài chính: nhiều cột Văn bản/Ngày hiệu lực xen giữa mỗi loại
thuế, các dòng Phần/Chương/Chú giải xen kẽ dòng mã HS).

Cách dùng:
    python scripts/import_tariff_bt2026.py --file "E:\\...\\BIEU THUE XNK 2026.xlsx" --sheet BT2026

Ghi đè/merge vào data/hs_tree.json (mặc định merge, dùng --replace để
nạp lại từ đầu).
"""
import sys
import io
import os
import re
import json
import argparse

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(BASE, "data")
TREE_PATH = os.path.join(DATA, "hs_tree.json")

# vị trí cột (0-based) đã xác nhận khớp thực tế với sheet BT2026
COL = {
    "level": 1,
    "code": 5,
    "desc_vn": 6,
    "desc_en": 7,
    "unit": 8,
    "thue_nk_thong_thuong": 10,
    "mfn": 13,
    "vat": 16,
    "chinh_sach": 99,
    "giam_vat": 100,
}
FTA_COLS = {
    "acfta": 19, "atiga": 22, "ajcep": 25, "vjepa": 28, "akfta": 31,
    "aanzfta": 34, "aifta": 37, "vkfta": 40, "vcfta": 43, "vneaeu": 46,
    "cptpp": 49, "ahkfta": 52, "vncu": 55, "evfta": 58, "ukvfta": 61,
    "vnlao": 64, "vifta": 67, "rcept": 70,
}


def cell(row, idx):
    return row[idx] if idx < len(row) else None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--file", required=True)
    ap.add_argument("--sheet", default="BT2026")
    ap.add_argument("--replace", action="store_true")
    args = ap.parse_args()

    try:
        import openpyxl
    except ImportError:
        print("Cần cài đặt: pip install openpyxl")
        sys.exit(1)

    print(f"Đang mở file (có thể mất chút thời gian do file lớn)...")
    wb = openpyxl.load_workbook(args.file, read_only=True, data_only=True)
    if args.sheet not in wb.sheetnames:
        print(f"Không thấy sheet '{args.sheet}'. Các sheet có sẵn: {wb.sheetnames}")
        sys.exit(1)
    ws = wb[args.sheet]

    tree = {"_meta": {}, "codes": {}}
    if os.path.exists(TREE_PATH) and not args.replace:
        with open(TREE_PATH, "r", encoding="utf-8") as f:
            tree = json.load(f)
        tree.setdefault("codes", {})

    current_section = None
    current_chapter = None
    count = 0
    skipped_no_duty = 0

    # Ngăn xếp mô tả cha-con: level (int) -> desc tại level đó, được cập
    # nhật liên tục khi duyệt tuần tự; khi gặp dòng ở level N, mọi mô tả
    # đã lưu ở level > N coi như hết hiệu lực (đã sang nhánh khác).
    level_desc_vn = {}
    level_desc_en = {}

    def build_full_desc(level_map, leaf_text):
        try:
            lvl_int = int(level)
        except (TypeError, ValueError):
            lvl_int = None
        parts = []
        if lvl_int is not None:
            for lv in sorted(k for k in level_map if k < lvl_int):
                parts.append(level_map[lv])
        if leaf_text:
            parts.append(leaf_text)
        return " - ".join(p for p in parts if p)

    for row in ws.iter_rows(values_only=True):
        level = cell(row, COL["level"])
        raw_code = cell(row, COL["code"])
        desc_vn = cell(row, COL["desc_vn"])
        desc_en = cell(row, COL["desc_en"])

        if level == "Phần" and desc_vn:
            current_section = str(desc_vn).strip()
            continue
        if level == "Chương" and desc_vn:
            current_chapter = str(desc_vn).strip()
            continue

        # Cập nhật ngăn xếp mô tả cha-con cho mọi dòng có level số nguyên
        # (kể cả dòng không có mã riêng, ví dụ dòng gộp nhóm trung gian
        # "- - Có mặt cắt ngang hình tròn:" không kèm mã 8 số).
        if isinstance(level, int) and desc_vn:
            # xóa các level sâu hơn (đã rẽ sang nhánh cha mới)
            for lv in list(level_desc_vn):
                if lv >= level:
                    del level_desc_vn[lv]
            for lv in list(level_desc_en):
                if lv >= level:
                    del level_desc_en[lv]
            level_desc_vn[level] = str(desc_vn).strip()
            level_desc_en[level] = str(desc_en).strip() if desc_en else ""

        if not raw_code:
            continue
        code = re.sub(r"\D", "", str(raw_code))
        if not code:
            continue

        thue_nk = cell(row, COL["thue_nk_thong_thuong"])
        mfn = cell(row, COL["mfn"])
        if thue_nk is None and mfn is None:
            # dòng nhóm/phân nhóm trung gian, không phải dòng có thuế suất thật
            skipped_no_duty += 1
            continue

        full_desc_vn = build_full_desc(level_desc_vn, str(desc_vn).strip() if desc_vn else "")
        full_desc_en = build_full_desc(level_desc_en, str(desc_en).strip() if desc_en else "")

        entry = {
            "desc_vn": full_desc_vn,
            "desc_en": full_desc_en,
            "section": current_section or "",
            "chapter": current_chapter or "",
            "unit": str(cell(row, COL["unit"]) or "").strip(),
            "thue_nk_thong_thuong": str(thue_nk).strip() if thue_nk is not None else "",
            "mfn": str(mfn).strip() if mfn is not None else "",
            "vat": str(cell(row, COL["vat"]) or "").strip(),
            "chinh_sach": str(cell(row, COL["chinh_sach"]) or "").strip(),
        }
        fta = {}
        for key, idx in FTA_COLS.items():
            v = cell(row, idx)
            if v is not None and str(v).strip() != "":
                fta[key] = str(v).strip()
        if fta:
            entry["fta"] = fta

        tree["codes"][code] = entry
        count += 1
        if count % 2000 == 0:
            print(f"  ... đã xử lý {count} mã")

    with open(TREE_PATH, "w", encoding="utf-8") as f:
        json.dump(tree, f, ensure_ascii=False, indent=2)

    print(f"Hoàn tất. Đã nhập {count} mã HS (bỏ qua {skipped_no_duty} dòng nhóm/phân nhóm trung gian).")
    print(f"Ghi vào: {TREE_PATH}")


if __name__ == "__main__":
    main()
